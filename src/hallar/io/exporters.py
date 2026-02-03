"""Export utilities for Hallar DSS results."""

import json
from pathlib import Path
from typing import Dict, Any, List, Optional, TYPE_CHECKING
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

if TYPE_CHECKING:
    from hallar.simulation.monte_carlo import FullAnalysisResults


def export_to_json(
    results: "FullAnalysisResults",
    output_path: str | Path,
    indent: int = 2,
) -> None:
    """
    Export analysis results to JSON format.

    Args:
        results: FullAnalysisResults object from run_full_analysis
        output_path: Path to output JSON file
    """
    output = results.to_dict()
    output["export_timestamp"] = datetime.now().isoformat()

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=indent, default=str)


def export_to_excel(
    results: "FullAnalysisResults",
    output_path: str | Path,
) -> None:
    """
    Export analysis results to Excel format.

    Args:
        results: FullAnalysisResults object from run_full_analysis
        output_path: Path to output Excel file
    """
    wb = Workbook()

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    good_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    bad_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Sheet 1: Executive Summary
    ws1 = wb.active
    ws1.title = "Executive Summary"

    ws1["A1"] = "HALLAR DECISION SUPPORT SYSTEM"
    ws1["A1"].font = Font(bold=True, size=16)
    ws1["A2"] = f"Analysis Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1["A3"] = f"Simulations: {results.n_simulations:,}"
    ws1["A4"] = f"Weight Profile: {results.weight_profile.name}"
    ws1["A5"] = f"Scenarios Evaluated: {len(results.scenario_results)}"
    ws1["A6"] = f"Passed Constraints: {sum(1 for r in results.scenario_results if r.passed_constraints)}"

    # Top scenarios table
    ws1["A8"] = "TOP SCENARIOS"
    ws1["A8"].font = Font(bold=True, size=14)

    headers = ["Rank", "ID", "Name", "Utility", "E[NPV]", "P(loss)", "Time (months)", "Constraints"]
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=9, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    for row_idx, result in enumerate(results.ranked_results[:10], 10):
        ws1.cell(row=row_idx, column=1, value=row_idx - 9).border = border
        ws1.cell(row=row_idx, column=2, value=result.scenario.id).border = border
        ws1.cell(row=row_idx, column=3, value=result.scenario.name).border = border
        ws1.cell(row=row_idx, column=4, value=round(result.utility, 3)).border = border
        ws1.cell(row=row_idx, column=5, value=round(result.metrics["expected_npv"], 0)).border = border
        ws1.cell(row=row_idx, column=6, value=f"{result.metrics['prob_loss']:.1%}").border = border
        ws1.cell(row=row_idx, column=7, value=round(result.metrics["expected_time"], 1)).border = border

        constraint_cell = ws1.cell(row=row_idx, column=8, value="PASS" if result.passed_constraints else "FAIL")
        constraint_cell.fill = good_fill if result.passed_constraints else bad_fill
        constraint_cell.border = border

    # Adjust column widths
    ws1.column_dimensions["A"].width = 8
    ws1.column_dimensions["B"].width = 10
    ws1.column_dimensions["C"].width = 35
    ws1.column_dimensions["D"].width = 10
    ws1.column_dimensions["E"].width = 12
    ws1.column_dimensions["F"].width = 10
    ws1.column_dimensions["G"].width = 15
    ws1.column_dimensions["H"].width = 12

    # Sheet 2: All Results
    ws2 = wb.create_sheet("All Results")

    # Create DataFrame of all results
    all_data = []
    for result in results.scenario_results:
        row = {
            "ID": result.scenario.id,
            "Name": result.scenario.name,
            "Utility": result.utility,
            "E[NPV]": result.metrics["expected_npv"],
            "NPV Std Dev": result.metrics["npv_std"],
            "NPV 10th %ile": result.metrics["npv_p10"],
            "NPV 90th %ile": result.metrics["npv_p90"],
            "CVaR 5%": result.metrics["cvar_05"],
            "P(loss)": result.metrics["prob_loss"],
            "E[Time] (months)": result.metrics["expected_time"],
            "Quality Score": result.metrics["quality_score"],
            "Affordability Score": result.metrics["affordability_score"],
            "Passed Constraints": "Yes" if result.passed_constraints else "No",
        }
        all_data.append(row)

    df = pd.DataFrame(all_data)

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws2.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border

    # Sheet 3: Constraint Details
    ws3 = wb.create_sheet("Constraint Details")

    ws3["A1"] = "Constraint Check Results"
    ws3["A1"].font = Font(bold=True, size=14)

    headers = ["Scenario", "Constraint", "Status", "Details"]
    for col, header in enumerate(headers, 1):
        cell = ws3.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    row = 4
    for result in results.scenario_results:
        if not result.passed_constraints:
            for violation_id, violation_msg in result.constraint_violations:
                ws3.cell(row=row, column=1, value=result.scenario.id).border = border
                ws3.cell(row=row, column=2, value=violation_id).border = border
                ws3.cell(row=row, column=3, value="FAIL").border = border
                ws3.cell(row=row, column=4, value=violation_msg).border = border
                ws3.cell(row=row, column=3).fill = bad_fill
                row += 1

    ws3.column_dimensions["A"].width = 15
    ws3.column_dimensions["B"].width = 25
    ws3.column_dimensions["C"].width = 10
    ws3.column_dimensions["D"].width = 60

    # Sheet 4: Weight Profiles
    ws4 = wb.create_sheet("Weight Analysis")

    ws4["A1"] = "Utility by Weight Profile"
    ws4["A1"].font = Font(bold=True, size=14)

    if hasattr(results, "weight_sensitivity") and results.weight_sensitivity:
        headers = ["Scenario"] + list(results.weight_sensitivity.keys())
        for col, header in enumerate(headers, 1):
            cell = ws4.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border

        for row_idx, result in enumerate(results.scenario_results, 4):
            ws4.cell(row=row_idx, column=1, value=result.scenario.id).border = border
            for col_idx, profile_name in enumerate(results.weight_sensitivity.keys(), 2):
                utility = results.weight_sensitivity[profile_name].get(result.scenario.id, 0)
                ws4.cell(row=row_idx, column=col_idx, value=round(utility, 3)).border = border

    wb.save(output_path)


def export_to_pdf(
    results: "FullAnalysisResults",
    output_path: str | Path,
) -> None:
    """
    Export analysis results to PDF format.

    Args:
        results: FullAnalysisResults object from run_full_analysis
        output_path: Path to output PDF file
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        PageBreak,
    )

    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle(
        "Title",
        parent=styles["Title"],
        fontSize=24,
        spaceAfter=30,
    )
    elements.append(Paragraph("Hallar Decision Support System", title_style))
    elements.append(Paragraph("Analysis Report", styles["Heading2"]))
    elements.append(Spacer(1, 20))

    # Summary info
    elements.append(Paragraph(f"<b>Analysis Date:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]))
    elements.append(Paragraph(f"<b>Simulations:</b> {results.n_simulations:,}", styles["Normal"]))
    elements.append(Paragraph(f"<b>Weight Profile:</b> {results.weight_profile.name}", styles["Normal"]))
    elements.append(Paragraph(f"<b>Scenarios Evaluated:</b> {len(results.scenario_results)}", styles["Normal"]))
    passed = sum(1 for r in results.scenario_results if r.passed_constraints)
    elements.append(Paragraph(f"<b>Passed Constraints:</b> {passed}", styles["Normal"]))
    elements.append(Spacer(1, 30))

    # Top scenarios table
    elements.append(Paragraph("Top Scenarios by Utility Score", styles["Heading2"]))
    elements.append(Spacer(1, 10))

    table_data = [["Rank", "ID", "Name", "Utility", "E[NPV]", "P(loss)", "Status"]]
    for rank, result in enumerate(results.ranked_results[:10], 1):
        table_data.append([
            str(rank),
            result.scenario.id,
            result.scenario.name[:25] + "..." if len(result.scenario.name) > 25 else result.scenario.name,
            f"{result.utility:.3f}",
            f"{result.metrics['expected_npv']:,.0f}",
            f"{result.metrics['prob_loss']:.1%}",
            "PASS" if result.passed_constraints else "FAIL",
        ])

    table = Table(table_data, colWidths=[0.5*inch, 0.6*inch, 2*inch, 0.7*inch, 1*inch, 0.7*inch, 0.6*inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
    ]))
    elements.append(table)
    elements.append(PageBreak())

    # Individual scenario summaries
    elements.append(Paragraph("Scenario Details", styles["Heading1"]))

    for result in results.ranked_results[:5]:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph(f"{result.scenario.id}: {result.scenario.name}", styles["Heading2"]))
        elements.append(Paragraph(result.scenario.description, styles["Normal"]))
        elements.append(Spacer(1, 10))

        metrics_data = [
            ["Metric", "Value"],
            ["Utility Score", f"{result.utility:.3f}"],
            ["Expected NPV", f"{result.metrics['expected_npv']:,.0f} M"],
            ["NPV Std Dev", f"{result.metrics['npv_std']:,.0f} M"],
            ["CVaR (5%)", f"{result.metrics['cvar_05']:,.0f} M"],
            ["P(loss)", f"{result.metrics['prob_loss']:.1%}"],
            ["Expected Time", f"{result.metrics['expected_time']:.1f} months"],
            ["Quality Score", f"{result.metrics['quality_score']:.1f}"],
            ["Affordability Score", f"{result.metrics['affordability_score']:.1f}"],
        ]

        metrics_table = Table(metrics_data, colWidths=[2*inch, 1.5*inch])
        metrics_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elements.append(metrics_table)

        # Constraint status
        if result.passed_constraints:
            elements.append(Paragraph("<font color='green'><b>All constraints passed</b></font>", styles["Normal"]))
        else:
            elements.append(Paragraph("<font color='red'><b>Constraint violations:</b></font>", styles["Normal"]))
            for vid, vmsg in result.constraint_violations:
                elements.append(Paragraph(f"  - {vmsg}", styles["Normal"]))

    doc.build(elements)
